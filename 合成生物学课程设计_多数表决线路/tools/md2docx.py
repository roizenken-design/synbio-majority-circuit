# -*- coding: utf-8 -*-
"""
Markdown -> DOCX 转换器（面向本课程设计书的子集）
支持: 标题/段落/加粗/行内代码/表格/图片/代码块/引用/分隔线/独立公式($$)
"""

import os
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else "docs/课程设计书.md"
DST = sys.argv[2] if len(sys.argv) > 2 else "docs/课程设计书.docx"

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.dirname(os.path.abspath(SRC))
MATH_DIR = os.path.join(BASE, "_math")
os.makedirs(MATH_DIR, exist_ok=True)

BODY_CN = "宋体"
BODY_EN = "Times New Roman"
HEAD_CN = "黑体"
CODE_FONT = "Consolas"


# ---------------------------------------------------------------- 基础工具
def set_run_font(run, cn=BODY_CN, en=BODY_EN, size=11, bold=False, italic=False,
                 color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)


def shade(cell_or_par, hexcolor):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    set_run_font(run, size=9, color=(0x66, 0x66, 0x66))


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline(par, text, size=11, cn=BODY_CN, en=BODY_EN, base_bold=False):
    """解析 **粗体**、`代码`、*斜体*"""
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = par.add_run(seg[2:-2])
            set_run_font(r, cn=cn, en=en, size=size, bold=True)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            set_run_font(r, cn=cn, en=CODE_FONT, size=size - 0.5,
                         color=(0x9B, 0x2C, 0x2C))
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            set_run_font(r, cn=cn, en=en, size=size, italic=True,
                         bold=base_bold)
        else:
            r = par.add_run(seg)
            set_run_font(r, cn=cn, en=en, size=size, bold=base_bold)


def render_math(tex, idx):
    """用 matplotlib mathtext 渲染独立公式为 PNG"""
    path = os.path.join(MATH_DIR, f"eq{idx}.png")
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${tex}$", fontsize=15, color="#1A202C")
    fig.savefig(path, dpi=400, bbox_inches="tight", pad_inches=0.06,
                facecolor="white")
    plt.close(fig)
    return path


# ---------------------------------------------------------------- 文档骨架
def build_cover(doc, title):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("课  程  设  计")
    set_run_font(r, cn=HEAD_CN, en=BODY_EN, size=20, bold=True)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, cn=HEAD_CN, en=BODY_EN, size=18, bold=True)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(40)

    for _ in range(4):
        doc.add_paragraph()

    for label in ["姓    名", "学    号", "专    业", "指导教师"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}：________________")
        set_run_font(r, cn=BODY_CN, en=BODY_EN, size=13)
        p.paragraph_format.space_after = Pt(10)

    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026 年 10 月")
    set_run_font(r, size=13)
    p.add_run().add_break(WD_BREAK.PAGE)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

    style = doc.styles["Normal"]
    style.font.name = BODY_EN
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    return doc


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    p.paragraph_format.keep_with_next = True
    size = {1: 16, 2: 13, 3: 11.5}[level]
    r = p.add_run(text)
    set_run_font(r, cn=HEAD_CN, en=BODY_EN, size=size, bold=True,
                 color=(0x1A, 0x20, 0x2C) if level == 1 else (0x2D, 0x37, 0x48))
    if level == 1:
        # 一级标题下加一条细线
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), "C9D2DB")
        pbdr.append(bottom)
        pPr.append(pbdr)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text)


def add_list_item(doc, text, marker):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.85)
    p.paragraph_format.first_line_indent = Cm(-0.85)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(marker + " ")
    set_run_font(r, size=11)
    add_inline(p, text, size=11)


def add_table(doc, rows):
    if not rows:
        return
    ncol = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        # 禁止单行跨页拆分
        trPr = table.rows[i]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(1)
            if i < len(rows) - 1:
                par.paragraph_format.keep_with_next = True
            txt = row[j] if j < len(row) else ""
            if i == 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = par.add_run(txt)
                set_run_font(r, cn=HEAD_CN, en=BODY_EN, size=9, bold=True)
                shade(cell, "E8EEF3")
            else:
                add_inline(par, txt, size=9)
    # 表后留白
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_together = True
    shade(p, "F5F7F9")
    for k, ln in enumerate(lines):
        r = p.add_run(ln)
        set_run_font(r, cn=BODY_CN, en=CODE_FONT, size=8.5,
                     color=(0x2D, 0x37, 0x48))
        if k < len(lines) - 1:
            r.add_break()


def add_image(doc, rel_path, caption):
    candidates = [
        os.path.join(ROOT, rel_path),
        os.path.join(BASE, rel_path),
    ]
    path = next((c for c in candidates if os.path.exists(c)), None)
    if path is None:
        print("  [warn] 图片缺失:", rel_path)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(15.2))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        r = cp.add_run(caption)
        set_run_font(r, size=9, color=(0x4A, 0x55, 0x68))


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.4
    add_inline(p, text, size=10.5)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------- 主解析
def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")

    # 提取封面标题
    title = "课程设计"
    start = 0
    for i, ln in enumerate(lines):
        if ln.startswith("# "):
            title = ln[2:].strip()
            start = i + 1
            break

    doc = setup_doc()
    build_cover(doc, title)

    i = start
    table_buf, code_buf = [], []
    in_code = False
    math_buf = None
    math_idx = 0
    first_hr_skipped = False

    while i < len(lines):
        raw = lines[i]
        ln = raw.rstrip()

        # 代码块
        if ln.startswith("```"):
            if in_code:
                add_code(doc, code_buf)
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        # 独立公式
        if ln.strip().startswith("$$"):
            body = ln.strip()
            if body == "$$":
                if math_buf is None:
                    math_buf = []
                else:
                    tex = " ".join(math_buf).strip()
                    math_buf = None
                    if tex:
                        math_idx += 1
                        img = render_math(tex, math_idx)
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(8)
                        p.add_run().add_picture(img, width=Cm(10))
                    i += 1
                    continue
            else:
                tex = body.strip("$").strip()
                if tex:
                    math_idx += 1
                    img = render_math(tex, math_idx)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(8)
                    p.add_run().add_picture(img, width=Cm(10))
                i += 1
                continue
            i += 1
            continue
        if math_buf is not None:
            math_buf.append(ln.strip())
            i += 1
            continue

        s = ln.strip()

        # 空行
        if not s:
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            # 表格结束则输出
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                add_table(doc, table_buf)
                table_buf = []
            continue

        # 分隔线
        if s == "---":
            if not first_hr_skipped:
                first_hr_skipped = True
            else:
                add_hr(doc)
            i += 1
            continue

        # 图片
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            add_image(doc, m.group(2), m.group(1))
            i += 1
            continue

        # 标题
        if s.startswith("### "):
            add_heading(doc, s[4:], 3)
            i += 1
            continue
        if s.startswith("## "):
            add_heading(doc, s[3:], 1)
            i += 1
            continue
        if s.startswith("# "):
            add_heading(doc, s[2:], 1)
            i += 1
            continue

        # 引用
        if s.startswith("> "):
            add_quote(doc, s[2:])
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            add_list_item(doc, m.group(2), m.group(1) + ".")
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            add_list_item(doc, m.group(1), "•")
            i += 1
            continue

        # 普通段落
        add_body(doc, s)
        i += 1

    if code_buf:
        add_code(doc, code_buf)

    doc.save(DST)
    print("saved:", DST)


if __name__ == "__main__":
    main()
